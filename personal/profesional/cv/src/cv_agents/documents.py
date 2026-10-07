"""DOCX structural validation and LibreOffice-backed PDF conversion.

Conversion intentionally invokes the external command ``soffice --headless
--convert-to pdf`` rather than assuming Microsoft Word. Install LibreOffice
and make ``soffice`` available on PATH, or pass an explicit command path. The
generated PDF still requires a final visual review.
"""

from __future__ import annotations

import shutil
import subprocess
from collections.abc import Callable, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader


class DocumentValidationError(ValueError):
    """Raised when a DOCX or PDF cannot be structurally validated."""


class PDFConversionError(RuntimeError):
    """Raised when LibreOffice cannot convert or validate a PDF output."""


@dataclass(frozen=True)
class DOCXValidation:
    path: Path
    paragraph_count: int
    table_count: int


@dataclass(frozen=True)
class PDFValidation:
    path: Path
    page_count: int
    extracted_text: str


def validate_docx(path: str | Path) -> DOCXValidation:
    """Confirm that a path is a readable DOCX package."""

    document_path = Path(path).expanduser().resolve()
    if not document_path.is_file():
        raise DocumentValidationError(f"DOCX file does not exist: {document_path}")
    try:
        document = Document(document_path)
    except Exception as exc:
        raise DocumentValidationError(
            f"DOCX file is not readable or is not a valid package: {document_path}: {exc}"
        ) from exc
    return DOCXValidation(
        path=document_path,
        paragraph_count=len(document.paragraphs),
        table_count=len(document.tables),
    )


def validate_pdf(path: str | Path, *, require_text: bool = False) -> PDFValidation:
    """Validate PDF page structure and extract text when the file supports it."""

    pdf_path = Path(path).expanduser().resolve()
    if not pdf_path.is_file():
        raise DocumentValidationError(f"PDF file does not exist: {pdf_path}")
    try:
        reader = PdfReader(str(pdf_path))
        page_count = len(reader.pages)
        extracted_text = "\n".join(page.extract_text() or "" for page in reader.pages).strip()
    except Exception as exc:
        raise DocumentValidationError(
            f"PDF file is not readable or is not a valid PDF: {pdf_path}: {exc}"
        ) from exc
    if page_count < 1:
        raise DocumentValidationError(f"PDF must contain at least one page: {pdf_path}")
    if require_text and not extracted_text:
        raise DocumentValidationError(f"PDF contains no extractable text: {pdf_path}")
    return PDFValidation(pdf_path, page_count, extracted_text)


def convert_docx_to_pdf(
    docx_path: str | Path,
    *,
    pdf_path: str | Path | None = None,
    soffice_command: str | Sequence[str] = "soffice",
    runner: Callable[..., Any] = subprocess.run,
) -> PDFValidation:
    """Convert a DOCX with LibreOffice and validate the resulting PDF."""

    source = validate_docx(docx_path).path
    output = (Path(pdf_path) if pdf_path is not None else source.with_suffix(".pdf")).expanduser()
    output = output.resolve()
    if output.suffix.lower() != ".pdf":
        raise PDFConversionError(f"PDF output must have a .pdf extension: {output}")
    output.parent.mkdir(parents=True, exist_ok=True)
    command_prefix = (
        [soffice_command] if isinstance(soffice_command, str) else list(soffice_command)
    )
    command = [
        *command_prefix,
        "--headless",
        "--convert-to",
        "pdf",
        "--outdir",
        str(output.parent),
        str(source),
    ]
    try:
        completed = runner(command, check=False, capture_output=True, text=True)
    except FileNotFoundError as exc:
        raise PDFConversionError(_soffice_remediation(soffice_command)) from exc
    except OSError as exc:
        raise PDFConversionError(
            f"Unable to execute LibreOffice command {command[0]!r}: {exc}"
        ) from exc

    returncode = getattr(completed, "returncode", 1)
    if returncode != 0:
        stderr = getattr(completed, "stderr", "") or getattr(completed, "stdout", "") or "no output"
        raise PDFConversionError(
            f"LibreOffice conversion failed with exit code {returncode}: {stderr}"
        )

    generated = output.parent / f"{source.stem}.pdf"
    if not generated.is_file():
        raise PDFConversionError(
            f"LibreOffice completed without creating {generated}. "
            "Check that the input DOCX is valid and the output directory is writable."
        )
    if generated != output:
        if output.exists():
            raise PDFConversionError(f"Refusing to overwrite existing PDF output: {output}")
        shutil.move(str(generated), str(output))
    return validate_pdf(output)


def _soffice_remediation(command: str | Sequence[str]) -> str:
    configured = command if isinstance(command, str) else " ".join(command)
    return (
        f"LibreOffice is required for DOCX-to-PDF conversion but {configured!r} was not found. "
        "Install LibreOffice and add soffice to PATH, or pass soffice_command with the "
        "absolute path to soffice."
    )
