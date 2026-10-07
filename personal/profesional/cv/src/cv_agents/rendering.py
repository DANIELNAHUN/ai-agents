"""Safe rendering of tailored content into the real DOCX output template.

The current ``files/templates/Template CV.docx`` contract requires each marker
below exactly once. The renderer changes marker text only, preserves all other
template content, and writes a separate output file. A final visual review is
still required for every template because DOCX layout can change with content.
"""

from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from .models import TailoredCVContent

REQUIRED_MARKERS = frozenset(
    {
        "{{resume_experience_fractalia}}",
        "{{list_bullet_goals_objetives_work_fractalia}}",
        "{{tools_list_fractalia}}",
        "{{resume_experience_cableclub}}",
        "{{list_bullet_goals_objetives_work_cableclub}}",
        "{{tools_list_cableclub}}",
    }
)
_MARKER_PATTERN = re.compile(r"\{\{[^{}]+\}\}")


class TemplateRenderError(ValueError):
    """Raised when a DOCX template is invalid or violates the marker contract."""


def render_tailored_cv(
    template_path: str | Path,
    output_path: str | Path,
    content: TailoredCVContent,
    *,
    overwrite: bool = False,
) -> Path:
    """Render ``content`` without modifying the source template."""

    template = _validate_template_path(template_path)
    output = Path(output_path).expanduser().resolve()
    if template == output:
        raise TemplateRenderError("Output path must be different from the source template")
    if output.exists() and not overwrite:
        raise TemplateRenderError(
            f"Refusing to overwrite existing output: {output}; pass overwrite=True explicitly"
        )
    if not isinstance(content, TailoredCVContent):
        content = TailoredCVContent.model_validate(content)

    try:
        document = Document(template)
    except Exception as exc:
        raise TemplateRenderError(f"Unable to read DOCX template {template}: {exc}") from exc

    paragraphs = list(_iter_paragraphs(document))
    occurrences = {marker: 0 for marker in REQUIRED_MARKERS}
    unknown_markers: set[str] = set()
    for paragraph in paragraphs:
        for marker in _MARKER_PATTERN.findall(paragraph.text):
            if marker in occurrences:
                occurrences[marker] += 1
            else:
                unknown_markers.add(marker)

    missing = sorted(marker for marker, count in occurrences.items() if count == 0)
    duplicated = sorted(marker for marker, count in occurrences.items() if count > 1)
    if missing or duplicated or unknown_markers:
        details = []
        if missing:
            details.append(f"missing required markers: {', '.join(missing)}")
        if duplicated:
            details.append(f"markers must occur once only: {', '.join(duplicated)}")
        if unknown_markers:
            details.append(f"unknown markers: {', '.join(sorted(unknown_markers))}")
        raise TemplateRenderError("Invalid DOCX marker contract; " + "; ".join(details))

    values: dict[str, str | list[str]] = {
        "{{resume_experience_fractalia}}": content.experience_fractalia.summary,
        "{{list_bullet_goals_objetives_work_fractalia}}": content.experience_fractalia.bullets,
        "{{tools_list_fractalia}}": ", ".join(content.experience_fractalia.tools),
        "{{resume_experience_cableclub}}": content.experience_cableclub.summary,
        "{{list_bullet_goals_objetives_work_cableclub}}": content.experience_cableclub.bullets,
        "{{tools_list_cableclub}}": ", ".join(content.experience_cableclub.tools),
    }

    for paragraph in paragraphs:
        marker_matches = _MARKER_PATTERN.findall(paragraph.text)
        if not marker_matches:
            continue
        marker = marker_matches[0]
        value = values[marker]
        if isinstance(value, list):
            _replace_bullet_paragraph(paragraph, marker, value)
        else:
            _replace_scalar_marker(paragraph, marker, value)

    output.parent.mkdir(parents=True, exist_ok=True)
    try:
        document.save(output)
    except Exception as exc:
        raise TemplateRenderError(f"Unable to write rendered DOCX {output}: {exc}") from exc
    return output


def _validate_template_path(path: str | Path) -> Path:
    template = Path(path).expanduser().resolve()
    if not template.is_file():
        raise TemplateRenderError(f"DOCX template does not exist: {template}")
    if template.suffix.lower() != ".docx":
        raise TemplateRenderError(f"DOCX template must have a .docx extension: {template}")
    return template


def _iter_paragraphs(document: Document) -> Iterable[Paragraph]:
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _replace_scalar_marker(paragraph: Paragraph, marker: str, value: str) -> None:
    matching_runs = [run for run in paragraph.runs if marker in run.text]
    if len(matching_runs) != 1:
        raise TemplateRenderError(
            f"Marker {marker} must be fully contained in exactly one DOCX run"
        )
    matching_runs[0].text = matching_runs[0].text.replace(marker, value)


def _replace_bullet_paragraph(paragraph: Paragraph, marker: str, bullets: list[str]) -> None:
    if paragraph.text != marker:
        raise TemplateRenderError(
            f"Bullet marker {marker} must occupy its whole paragraph in the DOCX template"
        )
    matching_runs = [run for run in paragraph.runs if marker in run.text]
    if len(matching_runs) != 1:
        raise TemplateRenderError(
            f"Marker {marker} must be fully contained in exactly one DOCX run"
        )

    source_paragraph = paragraph._p
    parent = source_paragraph.getparent()
    insertion_index = parent.index(source_paragraph)
    paragraph_properties = source_paragraph.find(qn("w:pPr"))
    run_properties = matching_runs[0]._r.find(qn("w:rPr"))
    lines = [f"• {bullet}" for bullet in bullets] or [""]

    for offset, line in enumerate(lines):
        new_paragraph = _build_paragraph(line, paragraph_properties, run_properties)
        if offset == 0:
            parent.replace(source_paragraph, new_paragraph)
        else:
            parent.insert(insertion_index + offset, new_paragraph)


def _build_paragraph(line: str, paragraph_properties: object, run_properties: object) -> object:
    paragraph = OxmlElement("w:p")
    if paragraph_properties is not None:
        paragraph.append(deepcopy(paragraph_properties))
    run = OxmlElement("w:r")
    if run_properties is not None:
        run.append(deepcopy(run_properties))
    text = OxmlElement("w:t")
    if line[:1].isspace() or line[-1:].isspace():
        text.set(qn("xml:space"), "preserve")
    text.text = line
    run.append(text)
    paragraph.append(run)
    return paragraph
