from pathlib import Path
from typing import Any

from docx import Document

from cv_agents.documents import PDFValidation
from cv_agents.models import GapAnalysis, JobRequirements
from cv_agents.pipeline import GeneratedDocumentPaths, generate_application_documents
from cv_agents.repository import ApplicationRecord


def _template(path: Path) -> None:
    document = Document()
    for marker in (
        "{{resume_experience_fractalia}}",
        "{{list_bullet_goals_objetives_work_fractalia}}",
        "{{tools_list_fractalia}}",
        "{{resume_experience_cableclub}}",
        "{{list_bullet_goals_objetives_work_cableclub}}",
        "{{tools_list_cableclub}}",
    ):
        paragraph = document.add_paragraph()
        if marker.startswith("{{tools"):
            paragraph.add_run("Tools: ")
        paragraph.add_run(marker)
    document.save(path)


def test_generation_service_consumes_record_and_keeps_fast_path_separate(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    _template(template)
    record = ApplicationRecord(
        application_id="app-synthetic",
        source="Synthetic job",
        job_text="Build APIs.",
        requirements=JobRequirements(summary="Build APIs."),
        profile_snapshot="# Profile\n\nPython experience.",
        gap_analysis=GapAnalysis(fit_score=70),
        status="analyzed",
        created_at="2026-01-01T00:00:00+00:00",
        updated_at="2026-01-01T00:00:00+00:00",
    )

    def invoke(_: dict[str, Any]) -> dict[str, Any]:
        return {
                "structured_response": {
                "experience_fractalia": {
                    "summary": "Python automation",
                    "bullets": [],
                    "tools": [],
                },
                "experience_cableclub": {
                    "summary": "Systems development",
                    "bullets": [],
                    "tools": [],
                },
            }
        }

    def fake_converter(docx_path: Path, *, pdf_path: Path) -> PDFValidation:
        pdf_path.write_bytes(b"synthetic pdf boundary")
        return PDFValidation(pdf_path.resolve(), 1, "")

    result = generate_application_documents(
        record,
        template_path=template,
        output_dir=tmp_path / "generated",
        generator_invoke=invoke,
        converter=fake_converter,
    )

    assert isinstance(result, GeneratedDocumentPaths)
    assert result.docx_path.is_file()
    assert result.pdf_path == (tmp_path / "generated" / "app-synthetic.pdf").resolve()
