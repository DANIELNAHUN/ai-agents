from pathlib import Path

import pytest
from docx import Document

from cv_agents.models import TailoredCVContent, TailoredExperience
from cv_agents.rendering import TemplateRenderError, render_tailored_cv


def _content() -> TailoredCVContent:
    return TailoredCVContent(
        experience_fractalia=TailoredExperience(
            summary="Built Python automations.",
            bullets=["Created API integrations.", "Added test coverage."],
            tools=["Python", "Pytest"],
        ),
        experience_cableclub=TailoredExperience(
            summary="Led internal systems development.",
            bullets=["Delivered operational modules."],
            tools=["Docker"],
        ),
    )


def _template(path: Path, *, omit: str | None = None) -> None:
    document = Document()
    document.add_paragraph("STATIC TEMPLATE CONTENT")
    markers = [
        "{{resume_experience_fractalia}}",
        "{{list_bullet_goals_objetives_work_fractalia}}",
        "{{tools_list_fractalia}}",
        "{{resume_experience_cableclub}}",
        "{{list_bullet_goals_objetives_work_cableclub}}",
        "{{tools_list_cableclub}}",
    ]
    for marker in markers:
        if marker == omit:
            continue
        paragraph = document.add_paragraph()
        if marker.startswith("{{tools"):
            paragraph.add_run("Tools: ")
        paragraph.add_run(marker)
    document.save(path)


def test_renderer_maps_markers_and_preserves_source_template(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    output = tmp_path / "generated" / "application.docx"
    _template(template)
    original = template.read_bytes()

    result = render_tailored_cv(template, output, _content())

    assert result == output.resolve()
    assert template.read_bytes() == original
    rendered = Document(output)
    texts = [paragraph.text for paragraph in rendered.paragraphs]
    assert "STATIC TEMPLATE CONTENT" in texts
    assert "Built Python automations." in texts
    assert "• Created API integrations." in texts
    assert "• Added test coverage." in texts
    assert "Tools: Python, Pytest" in texts
    assert all("{{" not in text for text in texts)


def test_renderer_reports_missing_required_marker(tmp_path: Path) -> None:
    template = tmp_path / "template.docx"
    _template(template, omit="{{tools_list_cableclub}}")

    with pytest.raises(TemplateRenderError, match="missing required markers"):
        render_tailored_cv(template, tmp_path / "output.docx", _content())
