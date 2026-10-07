from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter

from cv_agents.documents import (
    DocumentValidationError,
    PDFConversionError,
    convert_docx_to_pdf,
    validate_docx,
    validate_pdf,
)


def test_docx_and_pdf_validators_report_structure(tmp_path: Path) -> None:
    docx_path = tmp_path / "input.docx"
    Document().save(docx_path)
    pdf_path = tmp_path / "input.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with pdf_path.open("wb") as handle:
        writer.write(handle)

    docx_result = validate_docx(docx_path)
    pdf_result = validate_pdf(pdf_path)

    assert docx_result.path == docx_path.resolve()
    assert pdf_result.page_count == 1
    assert pdf_result.extracted_text == ""


def test_converter_invokes_soffice_and_validates_pdf(tmp_path: Path) -> None:
    docx_path = tmp_path / "application.docx"
    Document().save(docx_path)
    calls: list[list[str]] = []

    def fake_runner(command: list[str], **_: object) -> object:
        calls.append(command)
        writer = PdfWriter()
        writer.add_blank_page(width=612, height=792)
        with (tmp_path / "application.pdf").open("wb") as handle:
            writer.write(handle)
        return type("Completed", (), {"returncode": 0, "stderr": ""})()

    result = convert_docx_to_pdf(docx_path, runner=fake_runner)

    assert result.path == (tmp_path / "application.pdf").resolve()
    assert calls[0][:4] == ["soffice", "--headless", "--convert-to", "pdf"]


def test_converter_reports_exact_libreoffice_remediation(tmp_path: Path) -> None:
    docx_path = tmp_path / "application.docx"
    Document().save(docx_path)

    def missing_runner(*_: object, **__: object) -> object:
        raise FileNotFoundError("soffice")

    with pytest.raises(PDFConversionError, match="Install LibreOffice.*soffice"):
        convert_docx_to_pdf(docx_path, runner=missing_runner)


def test_validators_reject_unreadable_files(tmp_path: Path) -> None:
    invalid = tmp_path / "invalid.docx"
    invalid.write_text("not a docx", encoding="utf-8")

    with pytest.raises(DocumentValidationError, match="not a valid package"):
        validate_docx(invalid)
